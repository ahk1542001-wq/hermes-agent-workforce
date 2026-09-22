"""Truth-subset validation and ATS-readable DOCX/PDF/TXT rendering."""

from __future__ import annotations

import json
import os
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4  # type: ignore[import-untyped]
from reportlab.pdfbase.pdfmetrics import stringWidth  # type: ignore[import-untyped]
from reportlab.pdfgen.canvas import Canvas  # type: ignore[import-untyped]

from .models import CandidateFact

_HEADINGS = (
    "Professional Summary",
    "Professional Experience",
    "Selected AI Projects",
    "Technical Skills",
    "Selected Credentials",
    "Education",
    "Languages",
)
_BOLD_TERMS = (
    "Model Context Protocol (MCP)",
    "Google Cloud Run",
    "Google ADK",
    "Google Sheets",
    "human-in-the-loop",
    "least-privilege IAM",
    "FastAPI",
    "Python",
    "Docker",
    "Qdrant",
    "OpenAI",
    "Anthropic",
    "Groq",
    "Ollama",
    "Telegram",
    "Remotion",
    "ClickHouse",
    "Vertex AI",
    "BigQuery",
    "n8n",
    "RAG",
    "API",
    "LLM",
)
_BOLD_PATTERN = re.compile(
    "(" + "|".join(re.escape(term) for term in sorted(_BOLD_TERMS, key=len, reverse=True)) + ")",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class DraftClaim:
    text: str
    fact_id: str


@dataclass(frozen=True)
class ApplicationPackDraft:
    identity: DraftClaim
    headline: DraftClaim
    summary: DraftClaim
    skills: tuple[DraftClaim, ...]
    experience: tuple[DraftClaim, ...]
    education: tuple[DraftClaim, ...]
    languages: tuple[DraftClaim, ...] = ()
    contact: tuple[DraftClaim, ...] = ()
    projects: tuple[DraftClaim, ...] = ()
    credentials: tuple[DraftClaim, ...] = ()
    hidden_text: tuple[str, ...] = ()
    keyword_stuffing: tuple[str, ...] = ()


@dataclass(frozen=True)
class TruthReport:
    supported: bool
    unsupported_claims: tuple[str, ...]
    reason_codes: tuple[str, ...]


def _claims(draft: ApplicationPackDraft) -> tuple[DraftClaim, ...]:
    return (
        draft.identity,
        draft.headline,
        draft.summary,
        *draft.skills,
        *draft.experience,
        *draft.education,
        *draft.languages,
        *draft.contact,
        *draft.projects,
        *draft.credentials,
    )


def validate_truth_subset(
    draft: ApplicationPackDraft,
    facts: list[CandidateFact],
) -> TruthReport:
    """Require exact owner-approved wording after whitespace/case normalization."""

    approved = {fact.fact_id: fact for fact in facts if fact.verified}
    unsupported: list[str] = []
    reasons: list[str] = []
    seen_text: set[str] = set()
    if any(value.strip() for value in draft.hidden_text):
        reasons.append("HIDDEN_TEXT")
    if any(value.strip() for value in draft.keyword_stuffing):
        reasons.append("KEYWORD_STUFFING")

    for claim in _claims(draft):
        normalized = " ".join(claim.text.split()).casefold()
        fact = approved.get(claim.fact_id)
        if fact is None:
            unsupported.append(claim.text)
            if "UNAPPROVED_FACT" not in reasons:
                reasons.append("UNAPPROVED_FACT")
            continue
        allowed_normalized = " ".join(fact.allowed_wording.split()).casefold()
        if not normalized or normalized != allowed_normalized:
            unsupported.append(claim.text)
            if "UNSUPPORTED_CLAIM" not in reasons:
                reasons.append("UNSUPPORTED_CLAIM")
        if normalized in seen_text:
            if "KEYWORD_STUFFING" not in reasons:
                reasons.append("KEYWORD_STUFFING")
        seen_text.add(normalized)

    return TruthReport(
        supported=not unsupported and not reasons,
        unsupported_claims=tuple(unsupported),
        reason_codes=tuple(reasons),
    )


def _validated_lines(
    draft: ApplicationPackDraft,
    facts: list[CandidateFact],
) -> list[str]:
    report = validate_truth_subset(draft, facts)
    if not report.supported:
        raise ValueError(f"draft contains unsupported content: {','.join(report.reason_codes)}")
    lines = [draft.identity.text]
    if draft.contact:
        lines.append(" | ".join(claim.text for claim in draft.contact))
    lines.extend(
        [
            draft.headline.text,
            "Professional Summary",
            draft.summary.text,
            "Professional Experience",
            *(claim.text for claim in draft.experience),
        ]
    )
    if draft.projects:
        lines.extend(["Selected AI Projects", *(claim.text for claim in draft.projects)])
    if draft.credentials:
        lines.extend(
            ["Selected Credentials", " | ".join(claim.text for claim in draft.credentials)]
        )
    lines.extend(
        [
            "Technical Skills",
            *(claim.text for claim in draft.skills),
            "Education",
            *(claim.text for claim in draft.education),
        ]
    )
    if draft.languages:
        lines.extend(["Languages", " | ".join(claim.text for claim in draft.languages)])
    return lines


def _add_emphasized_text(paragraph: object, text: str) -> None:
    cursor = 0
    for match in _BOLD_PATTERN.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])  # type: ignore[attr-defined]
        run = paragraph.add_run(match.group(0))  # type: ignore[attr-defined]
        run.bold = True
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])  # type: ignore[attr-defined]


def _prepare_output(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists() and (path.is_symlink() or not path.is_file()):
        raise ValueError("document output must be a regular non-symlink file")


def _secure_file(path: Path) -> None:
    os.chmod(path, 0o600)


def render_resume_docx(
    draft: ApplicationPackDraft,
    facts: list[CandidateFact],
    output: Path,
) -> Path:
    _validated_lines(draft, facts)
    _prepare_output(output)
    document = Document()
    core = document.core_properties
    core.title = "ATS Resume"
    core.author = ""
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    normal_style = document.styles["Normal"]
    title_style = document.styles["Title"]
    heading_style = document.styles["Heading 1"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(11.5)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.space_after = Pt(2.4)
    normal_style.paragraph_format.line_spacing = 1.05
    title_style.font.name = "Times New Roman"
    title_style.font.size = Pt(18)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.space_after = Pt(1)
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading_style.font.name = "Times New Roman"
    heading_style.font.size = Pt(12)
    heading_style.font.bold = True
    heading_style.font.color.rgb = RGBColor(0, 0, 0)
    heading_style.paragraph_format.space_before = Pt(7)
    heading_style.paragraph_format.space_after = Pt(2)
    heading_style.paragraph_format.keep_with_next = True
    heading_properties = heading_style._element.get_or_add_pPr()
    heading_borders = heading_properties.find(qn("w:pBdr"))
    if heading_borders is None:
        heading_borders = OxmlElement("w:pBdr")
        heading_properties.append(heading_borders)
    bottom_border = heading_borders.find(qn("w:bottom"))
    if bottom_border is None:
        bottom_border = OxmlElement("w:bottom")
        heading_borders.append(bottom_border)
    bottom_border.set(qn("w:val"), "single")
    bottom_border.set(qn("w:sz"), "4")
    bottom_border.set(qn("w:space"), "1")
    bottom_border.set(qn("w:color"), "707070")
    if title_style._element.pPr is not None:
        border = title_style._element.pPr.find(qn("w:pBdr"))
        if border is not None:
            title_style._element.pPr.remove(border)
    bullet_style = document.styles["List Bullet"]
    bullet_style.font.name = "Times New Roman"
    bullet_style.font.size = Pt(11.5)
    bullet_style.font.color.rgb = RGBColor(0, 0, 0)
    bullet_style.paragraph_format.left_indent = Inches(0.2)
    bullet_style.paragraph_format.first_line_indent = Inches(-0.14)
    bullet_style.paragraph_format.space_after = Pt(1.8)
    bullet_style.paragraph_format.line_spacing = 1.04

    title = document.add_heading(draft.identity.text, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    if draft.contact:
        contact = document.add_paragraph(" | ".join(claim.text for claim in draft.contact))
        contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact.paragraph_format.space_after = Pt(1)
        for run in contact.runs:
            run.font.size = Pt(9.5)
    headline = document.add_paragraph()
    headline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    headline.paragraph_format.space_after = Pt(4)
    headline_run = headline.add_run(draft.headline.text)
    headline_run.italic = True
    headline_run.font.size = Pt(11)

    sections: list[tuple[str, tuple[DraftClaim, ...], bool]] = [
        ("Professional Summary", (draft.summary,), False),
        ("Professional Experience", draft.experience, True),
    ]
    if draft.projects:
        sections.append(("Selected AI Projects", draft.projects, True))
    sections.extend(
        [
            ("Technical Skills", draft.skills, True),
            ("Selected Credentials", draft.credentials, False),
            ("Education", draft.education, False),
            ("Languages", draft.languages, False),
        ]
    )
    for section_name, claims, use_bullets in sections:
        if not claims:
            continue
        heading = document.add_heading(section_name, level=1)
        heading.runs[0].font.color.rgb = RGBColor(0, 0, 0)
        if section_name in {"Languages", "Selected Credentials"}:
            document.add_paragraph(" | ".join(claim.text for claim in claims))
            continue
        for claim in claims:
            paragraph = document.add_paragraph(style="List Bullet" if use_bullets else None)
            _add_emphasized_text(paragraph, claim.text)
    document.save(str(output))
    _secure_file(output)
    return output


def render_resume_pdf(
    draft: ApplicationPackDraft,
    facts: list[CandidateFact],
    output: Path,
) -> Path:
    _validated_lines(draft, facts)
    _prepare_output(output)
    canvas = Canvas(str(output), pagesize=A4, pageCompression=1)
    canvas.setTitle("ATS Resume")
    width, height = A4
    margin = 44.0
    y = height - 38

    def wrap_line(value: str, font_name: str, font_size: float) -> list[str]:
        max_width = width - (2 * margin)
        wrapped: list[str] = []
        current = ""
        for word in value.split():
            candidate = f"{current} {word}".strip()
            if current and stringWidth(candidate, font_name, font_size) > max_width:
                wrapped.append(current)
                current = word
            else:
                current = candidate
        if current:
            wrapped.append(current)
        return wrapped or [""]

    def draw_centered(value: str, font_name: str, font_size: float, line_height: float) -> None:
        nonlocal y
        canvas.setFont(font_name, font_size)
        for physical_line in wrap_line(value, font_name, font_size):
            canvas.drawCentredString(width / 2, y, physical_line)
            y -= line_height

    def draw_section_heading(value: str) -> None:
        nonlocal y
        y -= 3
        canvas.setFont("Times-Bold", 12)
        canvas.drawString(margin, y, value)
        y -= 3
        canvas.setLineWidth(0.45)
        canvas.setStrokeColorRGB(0.35, 0.35, 0.35)
        canvas.line(margin, y, width - margin, y)
        canvas.setStrokeColorRGB(0, 0, 0)
        y -= 12

    def draw_body(value: str, *, bullet: bool = False) -> None:
        nonlocal y
        font_name = "Times-Roman"
        font_size = 11.5
        line_height = 13.6
        text_x = margin + (11 if bullet else 0)
        max_text_width = width - margin - text_x
        wrapped: list[str] = []
        current = ""
        for word in value.split():
            candidate = f"{current} {word}".strip()
            if current and stringWidth(candidate, font_name, font_size) > max_text_width:
                wrapped.append(current)
                current = word
            else:
                current = candidate
        if current:
            wrapped.append(current)
        if bullet:
            canvas.circle(margin + 3, y + 3, 1.3, stroke=0, fill=1)
        canvas.setFont(font_name, font_size)
        for physical_line in wrapped or [""]:
            if y < margin:
                canvas.showPage()
                y = height - margin
                canvas.setFont(font_name, font_size)
            canvas.drawString(text_x, y, physical_line)
            y -= line_height
        y -= 1.8

    draw_centered(draft.identity.text, "Times-Bold", 18, 20)
    if draft.contact:
        draw_centered(
            " | ".join(claim.text for claim in draft.contact),
            "Times-Roman",
            9.5,
            11.5,
        )
    draw_centered(draft.headline.text, "Times-Italic", 11, 13.5)

    sections = [
        ("Professional Summary", (draft.summary,), False),
        ("Professional Experience", draft.experience, True),
    ]
    if draft.projects:
        sections.append(("Selected AI Projects", draft.projects, True))
    sections.extend(
        [
            ("Technical Skills", draft.skills, True),
            ("Selected Credentials", draft.credentials, False),
            ("Education", draft.education, False),
            ("Languages", draft.languages, False),
        ]
    )
    for section_name, claims, use_bullets in sections:
        if not claims:
            continue
        draw_section_heading(section_name)
        if section_name in {"Languages", "Selected Credentials"}:
            draw_body(" | ".join(claim.text for claim in claims))
            continue
        for claim in claims:
            draw_body(claim.text, bullet=use_bullets)
    canvas.save()
    _secure_file(output)
    return output


def extract_docx_text(path: Path) -> str:
    document = Document(str(path))
    return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())


def extract_pdf_text(path: Path) -> str:
    reader = PdfReader(path)
    return "\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()


def _normalize_core(value: str) -> str:
    return " ".join(value.split()).casefold()


def compare_core_content(left: str, right: str) -> bool:
    return _normalize_core(left) == _normalize_core(right)


def _synthetic_material() -> tuple[list[CandidateFact], ApplicationPackDraft]:
    values = {
        "identity": "Synthetic Candidate",
        "headline": "AI Automation Engineer",
        "summary": "Builds safe evidence backed automation workflows",
        "python": "Python",
        "n8n": "n8n",
        "experience": "Built tested automation workflows for small teams",
        "education": "Independent AI automation study",
        "contact-email": "candidate@example.com",
        "contact-location": "Bangkok, Thailand",
        "project": "Built an evidence-backed agent workflow with deterministic approval gates",
        "credential": ("Google AI Professional Certificate — Google / Coursera, issued March 2026"),
    }
    facts = [
        CandidateFact(
            fact_id=fact_id,
            category="synthetic",
            allowed_wording=text,
            source="synthetic fixture",
            verified=True,
        )
        for fact_id, text in values.items()
    ]
    draft = ApplicationPackDraft(
        identity=DraftClaim(values["identity"], "identity"),
        headline=DraftClaim(values["headline"], "headline"),
        summary=DraftClaim(values["summary"], "summary"),
        skills=(DraftClaim(values["python"], "python"), DraftClaim(values["n8n"], "n8n")),
        experience=(DraftClaim(values["experience"], "experience"),),
        education=(DraftClaim(values["education"], "education"),),
        contact=(
            DraftClaim(values["contact-email"], "contact-email"),
            DraftClaim(values["contact-location"], "contact-location"),
        ),
        projects=(DraftClaim(values["project"], "project"),),
        credentials=(DraftClaim(values["credential"], "credential"),),
    )
    return facts, draft


def render_synthetic_resume(output: Path) -> dict[str, Path]:
    created = not output.exists()
    output.mkdir(parents=True, exist_ok=True)
    if output.is_symlink() or not output.is_dir():
        raise ValueError("synthetic output must be a non-symlink directory")
    if created:
        os.chmod(output, 0o700)
    facts, draft = _synthetic_material()
    report = validate_truth_subset(draft, facts)
    docx_path = render_resume_docx(draft, facts, output / "resume.docx")
    pdf_path = render_resume_pdf(draft, facts, output / "resume.pdf")
    txt_path = output / "resume.txt"
    txt_path.write_text(extract_docx_text(docx_path), encoding="utf-8")
    _secure_file(txt_path)
    audit_path = output / "resume.audit.json"
    audit_path.write_text(
        json.dumps(
            {
                "synthetic": True,
                "fact_ids": [claim.fact_id for claim in _claims(draft)],
                "unsupported_claims": list(report.unsupported_claims),
            },
            indent=2,
            sort_keys=True,
        )
        + "\n",
        encoding="utf-8",
    )
    _secure_file(audit_path)
    return {"docx": docx_path, "pdf": pdf_path, "txt": txt_path, "audit": audit_path}


__all__ = [
    "ApplicationPackDraft",
    "DraftClaim",
    "TruthReport",
    "compare_core_content",
    "extract_docx_text",
    "extract_pdf_text",
    "render_resume_docx",
    "render_resume_pdf",
    "render_synthetic_resume",
    "validate_truth_subset",
]
