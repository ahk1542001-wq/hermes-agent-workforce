from __future__ import annotations

import json
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from pypdf import PdfReader
from reportlab.pdfbase.pdfmetrics import stringWidth

from hermes_job_scout.documents import (
    ApplicationPackDraft,
    DraftClaim,
    compare_core_content,
    extract_docx_text,
    extract_pdf_text,
    render_resume_docx,
    render_resume_pdf,
    render_synthetic_resume,
    validate_truth_subset,
)
from hermes_job_scout.models import CandidateFact


def _facts() -> list[CandidateFact]:
    return [
        CandidateFact(
            fact_id="identity",
            category="identity",
            allowed_wording="Synthetic Candidate",
            source="synthetic fixture",
            verified=True,
        ),
        CandidateFact(
            fact_id="headline",
            category="headline",
            allowed_wording="AI Automation Engineer",
            source="synthetic fixture",
            verified=True,
        ),
        CandidateFact(
            fact_id="summary",
            category="summary",
            allowed_wording="Builds safe evidence backed automation workflows",
            source="synthetic fixture",
            verified=True,
        ),
        CandidateFact(
            fact_id="python",
            category="skill",
            allowed_wording="Python",
            source="synthetic fixture",
            verified=True,
        ),
        CandidateFact(
            fact_id="n8n",
            category="skill",
            allowed_wording="n8n",
            source="synthetic fixture",
            verified=True,
        ),
        CandidateFact(
            fact_id="experience",
            category="experience",
            allowed_wording="Built tested automation workflows for small teams",
            source="synthetic fixture",
            verified=True,
        ),
        CandidateFact(
            fact_id="education",
            category="education",
            allowed_wording="Independent AI automation study",
            source="synthetic fixture",
            verified=True,
        ),
        CandidateFact(
            fact_id="thai",
            category="language",
            allowed_wording="Thai: A2 (beginner)",
            source="synthetic fixture",
            verified=True,
        ),
        CandidateFact(
            fact_id="contact-email",
            category="contact",
            allowed_wording="candidate@example.com",
            source="synthetic fixture",
            verified=True,
        ),
        CandidateFact(
            fact_id="contact-location",
            category="contact",
            allowed_wording="Bangkok, Thailand",
            source="synthetic fixture",
            verified=True,
        ),
        CandidateFact(
            fact_id="project",
            category="project",
            allowed_wording=(
                "Built an evidence-backed agent workflow with deterministic approval gates"
            ),
            source="synthetic fixture",
            verified=True,
        ),
        CandidateFact(
            fact_id="credential",
            category="credential",
            allowed_wording=(
                "Google AI Professional Certificate — Google / Coursera, issued March 2026"
            ),
            source="synthetic fixture",
            verified=True,
        ),
    ]


def _draft(**updates: object) -> ApplicationPackDraft:
    values: dict[str, object] = {
        "identity": DraftClaim("Synthetic Candidate", "identity"),
        "headline": DraftClaim("AI Automation Engineer", "headline"),
        "summary": DraftClaim("Builds safe evidence backed automation workflows", "summary"),
        "skills": (DraftClaim("Python", "python"), DraftClaim("n8n", "n8n")),
        "experience": (
            DraftClaim("Built tested automation workflows for small teams", "experience"),
        ),
        "education": (DraftClaim("Independent AI automation study", "education"),),
        "languages": (DraftClaim("Thai: A2 (beginner)", "thai"),),
        "contact": (
            DraftClaim("candidate@example.com", "contact-email"),
            DraftClaim("Bangkok, Thailand", "contact-location"),
        ),
        "projects": (
            DraftClaim(
                "Built an evidence-backed agent workflow with deterministic approval gates",
                "project",
            ),
        ),
        "credentials": (
            DraftClaim(
                "Google AI Professional Certificate — Google / Coursera, issued March 2026",
                "credential",
            ),
        ),
    }
    values.update(updates)
    return ApplicationPackDraft(**values)


def test_supported_exact_wording_with_normalized_case_and_whitespace_passes() -> None:
    normalized_variant = "  builds SAFE evidence backed automation workflows "
    draft = _draft(summary=DraftClaim(normalized_variant, "summary"))
    report = validate_truth_subset(draft, _facts())
    assert report.supported is True
    assert report.unsupported_claims == ()


@pytest.mark.parametrize(
    ("allowed", "claim"),
    [
        ("Beginner Python", "Python"),
        ("Builds safe evidence backed automation workflows", "safe automation workflows"),
        ("Built tested automation workflows for small teams", "automation workflows Built tested"),
    ],
)
def test_unapproved_shortening_or_reordering_is_rejected(allowed: str, claim: str) -> None:
    facts = _facts()
    facts.append(
        CandidateFact(
            fact_id="exact",
            category="experience",
            allowed_wording=allowed,
            source="synthetic fixture",
            verified=True,
        )
    )
    draft = _draft(experience=(DraftClaim(claim, "exact"),))
    report = validate_truth_subset(draft, facts)
    assert report.supported is False
    assert "UNSUPPORTED_CLAIM" in report.reason_codes


@pytest.mark.parametrize(
    ("updates", "reason"),
    [
        (
            {"experience": (DraftClaim("Led a global engineering team", "experience"),)},
            "UNSUPPORTED_CLAIM",
        ),
        ({"headline": DraftClaim("Senior AI Director", "headline")}, "UNSUPPORTED_CLAIM"),
        ({"skills": (DraftClaim("Kubernetes", "python"),)}, "UNSUPPORTED_CLAIM"),
        ({"hidden_text": ("Python n8n AI automation",)}, "HIDDEN_TEXT"),
        ({"keyword_stuffing": ("AI AI AI n8n n8n",)}, "KEYWORD_STUFFING"),
    ],
)
def test_invented_or_hidden_content_is_rejected(updates: dict[str, object], reason: str) -> None:
    report = validate_truth_subset(_draft(**updates), _facts())
    assert report.supported is False
    assert reason in report.reason_codes


def test_unsupported_metric_is_rejected() -> None:
    draft = _draft(
        experience=(DraftClaim("Improved automation revenue by 40 percent", "experience"),)
    )
    report = validate_truth_subset(draft, _facts())
    assert report.supported is False
    assert report.unsupported_claims == ("Improved automation revenue by 40 percent",)


@pytest.mark.parametrize(
    ("allowed", "claim"),
    [
        ("No production experience", "production experience"),
        ("Thai basic", "Thai"),
        ("Worked from June 2020 to December 2022", "Worked"),
    ],
)
def test_meaning_changing_omissions_are_rejected(allowed: str, claim: str) -> None:
    facts = _facts()
    facts.append(
        CandidateFact(
            fact_id="guarded",
            category="experience",
            allowed_wording=allowed,
            source="synthetic fixture",
            verified=True,
        )
    )
    draft = _draft(experience=(DraftClaim(claim, "guarded"),))

    report = validate_truth_subset(draft, facts)

    assert report.supported is False
    assert report.reason_codes == ("UNSUPPORTED_CLAIM",)


def test_repeated_supported_token_keyword_stuffing_is_rejected() -> None:
    draft = _draft(skills=(DraftClaim("Python Python Python", "python"),))
    report = validate_truth_subset(draft, _facts())
    assert report.supported is False
    assert report.unsupported_claims == ("Python Python Python",)


def test_unknown_or_unverified_fact_is_rejected() -> None:
    facts = _facts() + [
        CandidateFact(
            fact_id="unverified",
            category="skill",
            allowed_wording="Rust",
            source="synthetic fixture",
            verified=False,
        )
    ]
    draft = _draft(skills=(DraftClaim("Rust", "unverified"),))
    assert validate_truth_subset(draft, facts).reason_codes == ("UNAPPROVED_FACT",)


def test_rendered_docx_pdf_and_text_have_matching_core_content(tmp_path: Path) -> None:
    draft = _draft()
    docx_path = render_resume_docx(draft, _facts(), tmp_path / "resume.docx")
    pdf_path = render_resume_pdf(draft, _facts(), tmp_path / "resume.pdf")
    txt_path = tmp_path / "resume.txt"
    txt_path.write_text(extract_docx_text(docx_path), encoding="utf-8")

    docx_text = extract_docx_text(docx_path)
    pdf_text = extract_pdf_text(pdf_path)
    plain_text = txt_path.read_text(encoding="utf-8")

    for heading in (
        "Professional Summary",
        "Professional Experience",
        "Selected AI Projects",
        "Technical Skills",
        "Selected Credentials",
        "Education",
        "Languages",
    ):
        assert heading in docx_text
        assert heading in pdf_text
    assert "Thai: A2 (beginner)" in docx_text
    assert "candidate@example.com | Bangkok, Thailand" in docx_text
    assert docx_text.index("Professional Summary") < docx_text.index("Professional Experience")
    assert docx_text.index("Professional Experience") < docx_text.index("Selected AI Projects")
    assert docx_text.index("Selected AI Projects") < docx_text.index("Technical Skills")
    assert docx_text.index("Technical Skills") < docx_text.index("Selected Credentials")
    assert docx_text.index("Selected Credentials") < docx_text.index("Education")
    assert pdf_text.index("Professional Summary") < pdf_text.index("Professional Experience")
    assert pdf_text.index("Professional Experience") < pdf_text.index("Selected AI Projects")
    assert pdf_text.index("Selected AI Projects") < pdf_text.index("Technical Skills")
    assert pdf_text.index("Technical Skills") < pdf_text.index("Selected Credentials")
    assert pdf_text.index("Selected Credentials") < pdf_text.index("Education")
    assert compare_core_content(docx_text, pdf_text)
    assert compare_core_content(docx_text, plain_text)
    assert docx_text.strip() and pdf_text.strip()
    assert oct(docx_path.stat().st_mode & 0o777) == "0o600"
    assert oct(pdf_path.stat().st_mode & 0o777) == "0o600"

    with zipfile.ZipFile(docx_path) as archive:
        xml = archive.read("word/document.xml")
    assert b"<w:tbl>" not in xml
    assert b"<w:vanish" not in xml
    assert b'w:color w:val="FFFFFF"' not in xml


def test_docx_title_and_headings_render_black_without_title_rule(tmp_path: Path) -> None:
    path = render_resume_docx(_draft(), _facts(), tmp_path / "resume.docx")

    document = Document(path)
    title = document.paragraphs[0]
    headings = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text
        in {
            "Professional Summary",
            "Professional Experience",
            "Selected AI Projects",
            "Technical Skills",
            "Selected Credentials",
            "Education",
            "Languages",
        }
    ]

    assert title.style.name == "Title"
    title_style = document.styles["Title"]
    heading_style = document.styles["Heading 1"]
    assert title_style.font.color.rgb is not None
    assert str(title_style.font.color.rgb) == "000000"
    assert heading_style.font.color.rgb is not None
    assert str(heading_style.font.color.rgb) == "000000"
    assert title_style._element.pPr is None or title_style._element.pPr.find(qn("w:pBdr")) is None
    assert title.runs[0].font.color.rgb is not None
    assert str(title.runs[0].font.color.rgb) == "000000"
    assert title._p.pPr is None or title._p.pPr.find(qn("w:pBdr")) is None
    assert all(
        run.font.color.rgb is not None and str(run.font.color.rgb) == "000000"
        for heading in headings
        for run in heading.runs
    )
    assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert document.paragraphs[1].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert document.paragraphs[2].alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_docx_uses_compact_single_page_ats_layout_contract(tmp_path: Path) -> None:
    path = render_resume_docx(_draft(), _facts(), tmp_path / "resume.docx")

    document = Document(path)
    section = document.sections[0]
    normal = document.styles["Normal"]
    title = document.styles["Title"]
    heading = document.styles["Heading 1"]

    assert section.start_type is WD_SECTION.NEW_PAGE
    assert abs(section.page_width - Inches(8.27)) < 1_000
    assert abs(section.page_height - Inches(11.69)) < 1_000
    assert section.top_margin <= Inches(0.7)
    assert section.bottom_margin <= Inches(0.7)
    assert section.left_margin <= Inches(0.75)
    assert section.right_margin <= Inches(0.75)
    assert normal.font.name == "Times New Roman"
    assert normal.font.size == Pt(11.5)
    assert normal.paragraph_format.space_after <= Pt(3)
    assert title.font.name == "Times New Roman"
    assert title.font.size == Pt(18)
    assert title.paragraph_format.space_after <= Pt(4)
    assert heading.font.name == "Times New Roman"
    assert heading.font.size == Pt(12)
    assert heading.paragraph_format.space_before <= Pt(8)
    assert heading.paragraph_format.space_after <= Pt(2)
    assert heading.paragraph_format.keep_with_next is True
    heading_border = heading._element.pPr.find(qn("w:pBdr"))
    assert heading_border is not None
    bottom_border = heading_border.find(qn("w:bottom"))
    assert bottom_border is not None
    assert bottom_border.get(qn("w:val")) == "single"
    assert bottom_border.get(qn("w:color")) == "707070"
    bullet_paragraphs = [
        paragraph for paragraph in document.paragraphs if paragraph.style.name == "List Bullet"
    ]
    assert len(bullet_paragraphs) >= 4


def test_pdf_wraps_long_claims_inside_page_margins(tmp_path: Path) -> None:
    long_summary = (
        "Builds safe evidence backed automation workflows with deterministic approval "
        "controls, auditable decisions, and reliable human review before external action"
    )
    facts = _facts() + [
        CandidateFact(
            fact_id="long-summary",
            category="summary",
            allowed_wording=long_summary,
            source="synthetic fixture",
            verified=True,
        )
    ]
    path = render_resume_pdf(
        _draft(summary=DraftClaim(long_summary, "long-summary")),
        facts,
        tmp_path / "resume.pdf",
    )

    overflows: list[tuple[float, float, str]] = []
    page = PdfReader(path).pages[0]
    assert len(PdfReader(path).pages) == 1
    right_edge = float(page.mediabox.width) - 44

    def inspect_text(
        text: str,
        _cm: list[float],
        tm: list[float],
        font: dict[str, object] | None,
        font_size: float,
    ) -> None:
        value = text.strip()
        if not value:
            return
        base_font = str((font or {}).get("/BaseFont", "/Helvetica"))
        font_name = "Times-Bold" if "Bold" in base_font else "Times-Roman"
        rendered_right = float(tm[4]) + stringWidth(value, font_name, font_size)
        if rendered_right > right_edge + 0.5:
            overflows.append((rendered_right, right_edge, value))

    page.extract_text(visitor_text=inspect_text)

    assert overflows == []


def test_render_rejects_unsupported_claim_before_writing(tmp_path: Path) -> None:
    output = tmp_path / "resume.docx"
    with pytest.raises(ValueError, match="unsupported"):
        render_resume_docx(
            _draft(skills=(DraftClaim("Invented Skill", "python"),)),
            _facts(),
            output,
        )
    assert not output.exists()


def test_synthetic_renderer_writes_auditable_private_artifacts(tmp_path: Path) -> None:
    artifacts = render_synthetic_resume(tmp_path)
    assert set(artifacts) == {"docx", "pdf", "txt", "audit"}
    assert all(path.exists() for path in artifacts.values())
    audit = json.loads(artifacts["audit"].read_text(encoding="utf-8"))
    assert audit["synthetic"] is True
    assert audit["unsupported_claims"] == []
    assert compare_core_content(
        extract_docx_text(artifacts["docx"]), extract_pdf_text(artifacts["pdf"])
    )
